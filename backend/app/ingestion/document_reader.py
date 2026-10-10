from pathlib import Path
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser

import csv
import pymupdf as fitz
import openpyxl


@dataclass
class ExtractedDocument:
    text: str
    source_type: str
    source_filename: str
    segments: list[dict] = field(default_factory=list)


class HTMLTextExtractor(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skip_depth = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self.skip_depth += 1
        elif tag in ("p", "br", "div", "tr", "li"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self.skip_depth:
            self.skip_depth -= 1

    def handle_data(self, data):
        if not self.skip_depth:
            self.parts.append(data)

    def get_text(self):
        return " ".join(self.parts)


def read_text(path):
    text = path.read_text(encoding="utf-8-sig")

    return ExtractedDocument(
        text=text,
        source_type="text",
        source_filename=path.name,
        segments=[
            {
                "line": i,
                "text": value,
            }
            for i, value in enumerate(text.splitlines(), 1)
            if value.strip()
        ],
    )


def read_email(path):
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())

    subject = str(message.get("subject") or "")

    body_part = message.get_body(preferencelist=("plain",))

    if body_part:
        body = body_part.get_content()
    else:
        html_part = message.get_body(preferencelist=("html",))

        if html_part:
            parser = HTMLTextExtractor()
            parser.feed(html_part.get_content())
            body = parser.get_text()
        else:
            body = ""

    text = f"Subject: {subject}\n\n{body}"

    if not body.strip():
        raise ValueError("Email contains no readable body")

    return ExtractedDocument(
        text=text,
        source_type="email",
        source_filename=path.name,
        segments=[
            {
                "line": i,
                "text": value,
            }
            for i, value in enumerate(text.splitlines(), 1)
            if value.strip()
        ],
    )


def read_pdf(path):
    segments = []

    with fitz.open(path) as document:
        for page_number, page in enumerate(document, 1):
            page_text = page.get_text("text")

            if page_text.strip():
                segments.append(
                    {
                        "page": page_number,
                        "text": page_text,
                    }
                )

    if not segments:
        raise ValueError(
            "PDF contains no extractable text. " "Scanned PDFs require OCR."
        )

    text = "\n\n".join(
        f"[Page {segment['page']}]\n{segment['text']}" for segment in segments
    )

    return ExtractedDocument(
        text=text,
        source_type="pdf",
        source_filename=path.name,
        segments=segments,
    )


def format_row(headers, values):
    fields = []

    for header, value in zip(headers, values):
        if value is not None and str(value).strip():
            fields.append(f"{header}: {value}")

    return " | ".join(fields)


def read_excel(path):
    workbook = openpyxl.load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    segments = []

    try:
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows(values_only=True)

            try:
                header_row = next(rows)
            except StopIteration:
                continue

            headers = [
                str(value) if value is not None else f"Column {i}"
                for i, value in enumerate(header_row, 1)
            ]

            for row_number, values in enumerate(rows, 2):
                text = format_row(headers, values)

                if not text:
                    continue

                segments.append(
                    {
                        "sheet": sheet.title,
                        "row": row_number,
                        "text": text,
                    }
                )

    finally:
        workbook.close()

    if not segments:
        raise ValueError("No readable spreadsheet rows found")

    text = "\n".join(
        f"[Sheet {s['sheet']}, Row {s['row']}] {s['text']}" for s in segments
    )

    return ExtractedDocument(
        text=text,
        source_type="xlsx",
        source_filename=path.name,
        segments=segments,
    )


def read_csv(path):
    segments = []

    with path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:

        reader = csv.DictReader(file)

        if not reader.fieldnames:
            raise ValueError("CSV file has no headers")

        for row_number, row in enumerate(reader, 2):
            text = format_row(
                reader.fieldnames,
                [row.get(h) for h in reader.fieldnames],
            )

            if text:
                segments.append(
                    {
                        "row": row_number,
                        "text": text,
                    }
                )

    if not segments:
        raise ValueError("No readable CSV rows found")

    text = "\n".join(f"[Row {s['row']}] {s['text']}" for s in segments)

    return ExtractedDocument(
        text=text,
        source_type="csv",
        source_filename=path.name,
        segments=segments,
    )


def read_document(file_path):
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(path)

    extension = path.suffix.lower()

    readers = {
        ".txt": read_text,
        ".eml": read_email,
        ".pdf": read_pdf,
        ".xlsx": read_excel,
        ".csv": read_csv,
    }

    if extension not in readers:
        raise ValueError(f"Unsupported file format: {extension}")

    return readers[extension](path)
